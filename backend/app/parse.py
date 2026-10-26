"""Proposal parsing: text + figures from PDF / DOCX / TXT / MD, plus a heuristic brief for AI-offline mode."""
from __future__ import annotations

import io
import re
from collections import Counter
from pathlib import Path

from .config import settings
from .retrieval import tokenize
from .schemas import ProposalBrief

MIN_FIG_PX = 180


class ParseError(ValueError):
    pass


def parse_file(filename: str, data: bytes) -> tuple[str, list[bytes]]:
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise ParseError(f"File is larger than {settings.max_upload_mb} MB")
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext in (".txt", ".md", ".markdown", ""):
        return data.decode("utf-8", "ignore"), []
    raise ParseError("Supported formats: PDF, DOCX, TXT, MD")


def _pdf(data: bytes) -> tuple[str, list[bytes]]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ParseError(f"Could not read PDF: {e}") from e
    if doc.page_count > settings.max_pages:
        raise ParseError(f"PDF has {doc.page_count} pages; the limit is {settings.max_pages}")
    text, images, seen = [], [], set()
    for page in doc:
        text.append(page.get_text("text"))
        for img in page.get_images(full=True):
            if len(images) >= settings.max_images or img[0] in seen:
                continue
            seen.add(img[0])
            pix = pymupdf.Pixmap(doc, img[0])
            if pix.width < MIN_FIG_PX or pix.height < MIN_FIG_PX:
                continue
            if pix.n - pix.alpha >= 4:  # CMYK -> RGB
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
            images.append(pix.tobytes("png"))
    return "\n".join(text), images


def _docx(data: bytes) -> tuple[str, list[bytes]]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not read DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    images = []
    for rel in d.part.rels.values():
        if "image" in rel.reltype and len(images) < settings.max_images:
            blob = rel.target_part.blob
            images.append(_to_png(blob))
    return "\n".join(parts), [i for i in images if i]


def _to_png(blob: bytes) -> bytes | None:
    import pymupdf

    try:
        pix = pymupdf.Pixmap(blob)
        if pix.width < MIN_FIG_PX or pix.height < MIN_FIG_PX:
            return None
        if pix.n - pix.alpha >= 4:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        return pix.tobytes("png")
    except Exception:
        return None


# ------------------------------------------------------------------ heuristic brief (used when Gemma is offline)
CUE = re.compile(r"\b(we (will|propose|build|develop|design)|our (system|app|solution|tool|platform)|feature|enables?|allows?|"
                 r"provides?|supports?|detects?|recommends?|generates?|offline|real[- ]time|multilingual)\b", re.I)


def _sentences(text: str) -> list[str]:
    text = re.sub(r"\s+", " ", text)
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if 25 < len(s.strip()) < 260]


def _section(text: str, *names: str) -> str:
    pat = re.compile(rf"^\s*(?:\d+[.)]?\s*)?(?:{'|'.join(names)})\s*:?\s*$", re.I | re.M)
    m = pat.search(text)
    if not m:
        return ""
    rest = text[m.end():]
    nxt = re.search(r"^\s*(?:\d+[.)]?\s*)?[A-Z][A-Za-z /&-]{2,40}:?\s*$", rest, re.M)
    return re.sub(r"\s+", " ", rest[: nxt.start()] if nxt else rest[:1200]).strip()[:600]


def heuristic_brief(text: str) -> ProposalBrief:
    lines = [re.sub(r"^#+\s*", "", l.strip()) for l in text.splitlines() if l.strip()]
    title = next((l for l in lines[:5] if 4 < len(l) < 120), "Untitled proposal")
    sents = _sentences(text)
    bullets = [re.sub(r"^[-*•\d.)\s]+", "", l) for l in lines if re.match(r"^\s*([-*•]|\d+[.)])\s+\S", l)]
    contributions = [b for b in bullets if 15 < len(b) < 200][:7]
    if len(contributions) < 3:
        contributions += [s for s in sents if CUE.search(s) and s not in contributions][: 7 - len(contributions)]
    toks = Counter(t for t in tokenize(text) if not t.isdigit() and len(t) > 3)
    return ProposalBrief(
        title=title,
        problem=_section(text, "problem", "problem statement", "motivation") or " ".join(sents[:2]),
        target_users=_section(text, "target users", "users", "audience") or "not stated",
        method=_section(text, "method", "methodology", "approach", "proposed solution", "solution") or "not stated",
        data=_section(text, "data", "dataset", "datasets") or "not stated",
        platform=_section(text, "platform", "tech stack", "technology", "technologies") or "not stated",
        expected_outcome=_section(text, "expected outcome", "outcomes", "impact", "results") or "not stated",
        contributions=contributions[:7] or sents[:3],
        keywords=[w for w, _ in toks.most_common(10)],
        generated_by="heuristic",
    )
